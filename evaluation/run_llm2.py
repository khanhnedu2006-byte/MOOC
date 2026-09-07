"""Chạy OCR + LLM2 trên mọi ảnh, ghi thứ ĐỌC ĐƯỢC ra Excel.

    python -m evaluation.run_llm2                     # quét data/image
    python -m evaluation.run_llm2 --limit 10          # thử vài ảnh trước
    python -m evaluation.run_llm2 --images duong/dan  # thư mục khác

Bốn cột, không hơn:

    tên file ảnh | tên người nhận | tên chứng chỉ | thời gian

KHÔNG có cột kết luận, KHÔNG có cột lý do, KHÔNG so với Excel. Mục đích duy
nhất của file này là xem LLM2 ĐỌC RA GÌ, để đối chiếu bằng mắt với chính tấm
ảnh. Thêm cột kết luận vào đây là trộn hai câu hỏi khác nhau: "model đọc đúng
chữ trên ảnh chưa" và "chứng chỉ có hợp lệ không". Câu đầu chỉ tấm ảnh trả
lời được; câu sau cần dữ liệu eLIS.

Ảnh nào OCR hoặc LLM hỏng thì ba cột nội dung để TRỐNG, và tên file được liệt
kê lại ở cuối màn hình. Không nhét thông báo lỗi vào ô dữ liệu: ô trống nghĩa
là "không đọc được", trộn chữ lỗi vào đó sẽ làm hỏng khi lọc bằng Excel.
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

# PHẢI đặt TRƯỚC mọi import kéo theo thư viện tính toán (pypdfium2, Pillow,
# azure-*). Trên Windows dùng conda, hai bản libiomp5md.dll (OpenMP của Intel)
# bị nạp cùng lúc và chương trình chết ngay với "OMP: Error #15".
# Đặt sau import là VÔ TÁC DỤNG: DLL đã nạp xong trước khi dòng này chạy.
os.environ.setdefault("KMP_DUPLICATE_LIB_OK", "TRUE")

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

import file_utils                                      # noqa: E402
import llm_text                                        # noqa: E402
import ocr_azure                                       # noqa: E402

from evaluation import match_images                    # noqa: E402

FONT = "Arial"

COT = [("tên file ảnh", 54), ("tên người nhận", 28),
       ("tên chứng chỉ", 50), ("thời gian", 22)]


def run_all(paths, azure_client, limit=0) -> tuple[list[list], list[tuple[str, str]]]:
    """Chạy OCR + LLM2 cho từng ảnh.

    Trả về (dòng để ghi Excel, danh sách ảnh hỏng kèm lý do).
    """
    if limit:
        paths = paths[:limit]
    rows, failed = [], []
    for i, path in enumerate(paths, start=1):
        file_name = path.name
        print(f"  [{i}/{len(paths)}] {file_name[:56]} ... ", end="", flush=True)
        try:
            images = file_utils.read_as_images(path)
            ocr_text = ocr_azure.ocr_images(azure_client, images)
            t2 = llm_text.extract_from_text(ocr_text)
        except Exception as e:
            # Một ảnh hỏng KHÔNG được giết cả lượt chạy: những ảnh trước đã
            # tốn tiền Azure rồi.
            print(f"LỖI: {e}")
            failed.append((file_name, f"{type(e).__name__}: {e}"))
            rows.append([file_name, "", "", ""])
            continue

        rows.append([file_name, t2.recipient_name or "",
                     t2.certificate_name or "", t2.issue_date or ""])
        print(f"{(t2.recipient_name or '?')[:24]} | "
              f"{(t2.certificate_name or '?')[:30]} | {t2.issue_date or '?'}")
    return rows, failed


def write_excel(rows, path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "LLM2"

    for i, (name, rong) in enumerate(COT, start=1):
        o = ws.cell(row=1, column=i, value=name)
        o.font = Font(name=FONT, bold=True)
        o.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = rong

    for so, row in enumerate(rows, start=2):
        for i, gia_tri in enumerate(row, start=1):
            o = ws.cell(row=so, column=i, value=gia_tri)
            o.font = Font(name=FONT, size=10)
            o.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(
        description="Chạy OCR + LLM2 trên mọi ảnh, ghi thứ đọc được ra Excel")
    p.add_argument("--images", default=str(PROJECT_ROOT / "data" / "image"),
                   help="thư mục ảnh chứng chỉ (mặc định data/image)")
    p.add_argument("--out", default=str(PROJECT_ROOT / "evaluation" / "llm2_result.xlsx"),
                   help="file Excel ghi ra")
    p.add_argument("--limit", type=int, default=0, help="chỉ chạy N ảnh đầu")
    args = p.parse_args(argv)

    paths = match_images.list_images(args.images)
    print(f"{len(paths)} ảnh trong {args.images}")
    if args.limit:
        print(f"--limit {args.limit}: chỉ chạy {args.limit} ảnh đầu.")
    print()

    rows, failed = run_all(paths, ocr_azure.create_client(), args.limit)
    write_excel(rows, args.out)

    print(f"\nĐã ghi {args.out} — {len(rows)} dòng.")
    if failed:
        print(f"\n{len(failed)} ảnh KHÔNG đọc được (ba cột nội dung để trống):")
        for name, reason in failed:
            print(f"  {name}\n      {reason}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
