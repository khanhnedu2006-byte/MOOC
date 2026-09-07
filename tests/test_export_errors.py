"""Test bảng Excel gửi HR (test_export_errors).

Bảng này ĐI RA NGOÀI nhóm phát triển. Một dòng phân loại sai ở đây không chỉ
là bug — nó là một câu khẳng định sai về một nhân viên cụ thể, gửi cho HR.
Nên hàm phân tích nguyên nhân phải KHÔNG ĐOÁN BỪA: ca không quy được về mẫu
nào phải nói thẳng "cần mở ảnh kiểm tra".
"""

import pathlib
import sys

GOC = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(GOC))
sys.path.insert(0, str(GOC / "src"))

from evaluation import export_errors                   # noqa: E402


def _cap(reason, ai_ten="", ai_khoa="", ai_ngay="",
         elis_ten="Lê Hoàng Anh", elis_khoa="Claude Code in Action"):
    e = {"input_employee_name": elis_ten, "input_course_name": elis_khoa}
    r = {"reason": reason, "recipient_name": ai_ten,
         "certificate_name": ai_khoa, "issue_date": ai_ngay}
    return e, r


# ===== Ngày =====

def test_ngay_rong_va_ngay_rac_deu_la_KHONG_DOC_DUOC():
    """'' và '01/01/0001' là 'không đọc được', KHÁC với 'ngày ngoài kỳ'.

    Gộp hai thứ này chính là lỗi đang làm hệ thống chặn oan: nó lẫn lộn
    'không có bằng chứng' với 'bằng chứng cho thấy sai'.
    """
    for day in ("", "01/01/0001", "0001-01-01"):
        e, r = _cap("Ngày không hợp lệ", ai_ngay=day)
        assert "Không đọc được ngày" in export_errors.nguyen_nhan_tu_choi_oan(e, r)


def test_ngay_doc_duoc_nhung_ngoai_khoang_thi_noi_khac():
    e, r = _cap("Ngày không hợp lệ", ai_ngay="15/08/2025")
    ra = export_errors.nguyen_nhan_tu_choi_oan(e, r)
    assert "ngoài khoảng quy định" in ra
    assert "Không đọc được" not in ra


# ===== Tên =====

def test_bo_ten_dem_duoc_nhan_dien():
    """'Lê Hoàng Anh' -> chứng chỉ in 'Anh Le': tập từ con, không phải sai tên."""
    e, r = _cap("Tên không khớp", ai_ten="Anh Le")
    assert "bỏ bớt tên đệm" in export_errors.nguyen_nhan_tu_choi_oan(e, r)


def test_email_hoac_username_duoc_nhan_dien():
    for name in ("minhnt4487@gmail.com", "kieuhuuthanh23698"):
        e, r = _cap("Tên không khớp", ai_ten=name)
        assert "email/username" in export_errors.nguyen_nhan_tu_choi_oan(e, r)


def test_ten_khac_han_thi_KHONG_ket_luan_ma_bao_kiem_tra():
    """Ca quan trọng nhất: tên khác hẳn có thể là nộp nhầm — AI có thể ĐÚNG.

    Kết luận bừa 'hệ thống sai' ở đây là bỏ qua một ca gian lận thật.
    """
    e, r = _cap("Tên không khớp", ai_ten="ĐỖ VĂN KHẮC", elis_ten="Phạm Đình Tiến")
    ra = export_errors.nguyen_nhan_tu_choi_oan(e, r)
    assert "KHÁC HẲN" in ra and "kiểm tra" in ra


# ===== Tên khóa học =====

def test_ten_khoa_rut_gon_song_ngu():
    e, r = _cap("Tên khóa học không khớp",
                ai_khoa="BỘ QUY ĐỊNH CHÍNH SÁCH CẦN BIẾT FPT",
                elis_khoa="Bộ Quy định chính sách cần biết FPT - FPT Key "
                          "Regulations and Policies (English version)")
    assert "rút gọn" in export_errors.nguyen_nhan_tu_choi_oan(e, r)


def test_ten_khoa_lech_vai_tu():
    e, r = _cap("Tên khóa học không khớp",
                ai_khoa="Introduction to agent skills",
                elis_khoa="Intro to Agent Skills")
    assert "lệch vài từ" in export_errors.nguyen_nhan_tu_choi_oan(e, r)


def test_ten_khoa_khac_han_thi_bao_kiem_tra():
    e, r = _cap("Tên khóa học không khớp",
                ai_khoa="KỸ NĂNG BÁN HÀNG BẰNG CÂU CHUYỆN",
                elis_khoa="Storytelling trong bán hàng trên kênh cá nhân")
    assert "KHÁC HẲN" in export_errors.nguyen_nhan_tu_choi_oan(e, r)


def test_nhieu_truong_sai_thi_liet_ke_het():
    """Liệt kê MỌI nguyên nhân: sửa một cái mà ca vẫn trượt thì mất công vô ích."""
    e, r = _cap("Tên không khớp; Tên khóa học không khớp; Ngày không hợp lệ",
                ai_ten="Anh Le", ai_khoa="Claude Code in Action ver 2", ai_ngay="")
    ra = export_errors.nguyen_nhan_tu_choi_oan(e, r)
    assert ra.count(";") == 2


# ===== Dựng workbook =====

def _bo_du_lieu(tmp_path):
    import csv
    label = tmp_path / "eval_set.csv"
    with label.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["case_id", "image_path", "input_employee_name",
                    "input_course_name", "input_employee_code", "gt_verdict", "note"])
        w.writerow(["001", "data/image/a.jpg", "Lê Hoàng Anh", "Claude Code in Action",
                    "anhlh48", "APPROVED", "excel_row=2 | elis=APPROVED"])
        w.writerow(["002", "data/image/b.jpg", "Nguyễn Đồng Thiện", "SAN and NAS",
                    "thiennd8", "REJECTED",
                    "excel_row=3 | elis=REJECTED | lý do người duyệt: CB log trùng khóa"])
        w.writerow(["003", "data/image/c.jpg", "Trần Văn C", "ISO 27001",
                    "cvt1", "APPROVED", "excel_row=4 | elis=APPROVED"])
    ket_qua = tmp_path / "last_run.csv"
    with ket_qua.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["case_id", "image_path", "verdict", "stage", "reason", "error",
                    "recipient_name", "certificate_name", "issue_date"])
        w.writerow(["001", "", "REJECTED", "llm1_vs_llm2", "Tên không khớp", "",
                    "Anh Le", "Claude Code in Action", "16/03/2026"])
        w.writerow(["002", "", "APPROVED", "llm1", "Tên, khóa học và thời gian đều khớp",
                    "", "Nguyễn Đồng Thiện", "SAN and NAS", "31/03/2026"])
        w.writerow(["003", "", "APPROVED", "llm1", "khớp", "",
                    "Trần Văn C", "ISO 27001", "01/05/2026"])   # ĐÚNG -> không vào bảng
    return label, ket_qua


def test_tach_dung_hai_huong_va_bo_ca_dung(tmp_path):
    label, ket_qua = _bo_du_lieu(tmp_path)
    ra = tmp_path / "review.xlsx"

    duyet_oan, tu_choi_oan = export_errors.build(label, ket_qua, ra)

    assert (duyet_oan, tu_choi_oan) == (1, 1), "ca máy đoán ĐÚNG lọt vào bảng lệch"

    from openpyxl import load_workbook
    wb = load_workbook(ra)
    assert wb.sheetnames == ["Tổng quan", "Máy duyệt - Người từ chối",
                             "Máy từ chối - Người duyệt"]
    # Mỗi sheet: dòng 1 mô tả, dòng 2 tiêu đề, từ dòng 3 là dữ liệu.
    assert wb["Máy duyệt - Người từ chối"].cell(row=3, column=1).value == "002"
    assert wb["Máy từ chối - Người duyệt"].cell(row=3, column=1).value == "001"


def test_cot_hr_de_trong_cho_nguoi_dien(tmp_path):
    label, ket_qua = _bo_du_lieu(tmp_path)
    ra = tmp_path / "review.xlsx"
    export_errors.build(label, ket_qua, ra)

    from openpyxl import load_workbook
    ws = load_workbook(ra)["Máy từ chối - Người duyệt"]
    assert ws.cell(row=2, column=15).value == "Ý KIẾN HR"
    assert ws.cell(row=2, column=16).value == "KẾT LUẬN THỐNG NHẤT"
    assert ws.cell(row=3, column=15).value in (None, "")
    assert ws.cell(row=3, column=16).value in (None, "")


def test_ca_khong_chay_duoc_khong_vao_bang(tmp_path):
    """Ca hỏng kỹ thuật (verdict rỗng) không phải 'lệch với người duyệt'."""
    import csv
    label, ket_qua = _bo_du_lieu(tmp_path)
    with ket_qua.open("a", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerow(["004", "", "", "", "", "OSError: hỏng", "", "", ""])
    with label.open("a", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerow(["004", "data/image/d.jpg", "D", "K", "d1",
                                "APPROVED", ""])

    duyet_oan, tu_choi_oan = export_errors.build(label, ket_qua, tmp_path / "r.xlsx")
    assert (duyet_oan, tu_choi_oan) == (1, 1)
