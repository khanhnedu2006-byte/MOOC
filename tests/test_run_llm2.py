"""Test lệnh chạy OCR + LLM2 rồi ghi ra Excel (test_run_llm2).

File này CHỈ ghi thứ model đọc được, không kết luận gì. Ba tính chất phải
đúng, và cả ba đều hỏng lặng lẽ nếu sai:

  1. Chạy cho MỌI ảnh trong thư mục — bỏ sót ảnh nào là mất đúng mẫu cần soi.
  2. Ảnh hỏng để TRỐNG ô, không nhét chữ lỗi vào ô dữ liệu (hỏng bộ lọc Excel).
  3. Một ảnh hỏng không giết cả lượt chạy — ảnh trước đã tốn tiền Azure rồi.
"""

import pathlib
import sys

import pytest

GOC = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(GOC))
sys.path.insert(0, str(GOC / "src"))

from evaluation import run_llm2                        # noqa: E402
from schemas import ExtractedInfo                      # noqa: E402


def _doc(name="Bùi Đức Hòa", key="ISO 27001", day="01/08/2026"):
    return ExtractedInfo(recipient_name=name, certificate_name=key,
                         issue_date=day)


@pytest.fixture
def moi_truong(monkeypatch):
    monkeypatch.setattr(run_llm2.file_utils, "read_as_images", lambda p: [b"anh"])
    dem = {"azure": 0}

    def azure(_c, _i):
        dem["azure"] += 1
        return "text ocr"

    monkeypatch.setattr(run_llm2.ocr_azure, "ocr_images", azure)
    monkeypatch.setattr(run_llm2.llm_text, "extract_from_text", lambda t: _doc())
    return dem


def _anh(thu_muc, *name):
    thu_muc.mkdir(parents=True, exist_ok=True)
    for t in name:
        (thu_muc / t).write_bytes(b"x")
    return sorted(thu_muc.iterdir())


def test_chay_cho_MOI_anh_trong_thu_muc(moi_truong, tmp_path):
    paths = _anh(tmp_path / "anh", *[f"NV{i}_Khóa {i}.jpg" for i in range(7)])
    rows, failed = run_llm2.run_all(paths, None)
    assert len(rows) == 7 and moi_truong["azure"] == 7 and failed == []


def test_dung_4_COT_theo_dung_thu_tu(moi_truong, tmp_path):
    """Không có cột kết luận, không có cột lý do — đúng yêu cầu."""
    paths = _anh(tmp_path / "anh", "HOABD3_ISO 27001.jpg")
    rows, _ = run_llm2.run_all(paths, None)

    assert len(rows[0]) == 4
    assert rows[0] == ["HOABD3_ISO 27001.jpg", "Bùi Đức Hòa", "ISO 27001",
                       "01/08/2026"]

    ra = tmp_path / "kq.xlsx"
    run_llm2.write_excel(rows, ra)
    from openpyxl import load_workbook
    ws = load_workbook(ra).active
    assert [c.value for c in ws[1]] == [
        "tên file ảnh", "tên người nhận", "tên chứng chỉ", "thời gian"]
    assert ws.max_column == 4, "có cột thừa"


def test_truong_model_khong_doc_ra_thi_de_RONG(moi_truong, tmp_path, monkeypatch):
    """None -> ô trống, không phải chữ 'None'."""
    monkeypatch.setattr(run_llm2.llm_text, "extract_from_text",
                        lambda t: _doc(day=None))
    paths = _anh(tmp_path / "anh", "A_x.jpg")
    rows, _ = run_llm2.run_all(paths, None)
    assert rows[0][3] == ""


def test_anh_hong_de_TRONG_o_chu_khong_nhet_chu_loi(moi_truong, tmp_path, monkeypatch):
    """Nhét chữ lỗi vào ô dữ liệu là làm hỏng bộ lọc Excel của người đọc."""
    monkeypatch.setattr(run_llm2.ocr_azure, "ocr_images",
                        lambda c, i: (_ for _ in ()).throw(RuntimeError("Azure sập")))
    paths = _anh(tmp_path / "anh", "A_x.jpg")
    rows, failed = run_llm2.run_all(paths, None)

    assert rows[0] == ["A_x.jpg", "", "", ""]
    assert failed == [("A_x.jpg", "RuntimeError: Azure sập")]


def test_mot_anh_hong_khong_giet_ca_luot_chay(moi_truong, tmp_path, monkeypatch):
    attempt = {"n": 0}

    def thinh_thoang_hong(_c, _i):
        attempt["n"] += 1
        if attempt["n"] == 1:
            raise RuntimeError("hỏng")
        return "text ocr"

    monkeypatch.setattr(run_llm2.ocr_azure, "ocr_images", thinh_thoang_hong)
    paths = _anh(tmp_path / "anh", "A_x.jpg", "B_y.jpg")
    rows, failed = run_llm2.run_all(paths, None)

    assert len(rows) == 2 and len(failed) == 1
    assert rows[1][1] == "Bùi Đức Hòa", "ảnh sau không được chạy"


def test_limit_chan_truoc_khi_goi_azure(moi_truong, tmp_path):
    paths = _anh(tmp_path / "anh", *[f"NV{i}_K{i}.jpg" for i in range(10)])
    run_llm2.run_all(paths, None, limit=3)
    assert moi_truong["azure"] == 3
